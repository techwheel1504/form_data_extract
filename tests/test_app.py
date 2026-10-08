"""Exercise the real OpenAI SDK against a mocked HTTP transport; no paid calls."""
import base64
import csv
import io
import json
import os
import unittest
from unittest.mock import patch

try:
    import httpx2 as httpx
except ImportError:
    import httpx
import pymupdf
from openai import OpenAI
from openpyxl import load_workbook
from PIL import Image

from app import app
from extraction import document_content


def image_bytes(format='PNG'):
    output = io.BytesIO()
    Image.new('RGB', (400, 200), 'white').save(output, format)
    return output.getvalue()


class AppTests(unittest.TestCase):
    def setUp(self):
        self.client = app.test_client()
        self.env = patch.dict(os.environ, {'OPENAI_API_KEY': 'test-secret-not-a-real-key', 'OPENAI_MODEL': 'gpt-4.1-mini'})
        self.env.start()
        self.addCleanup(self.env.stop)
        self.calls = []
        self.status_code = 200
        self.response_status = 'completed'
        self.refusal = False
        self.fields = [{'label': 'Name', 'value': 'Alice Smith'}, {'label': 'Phone', 'value': ''}]
        self.factory = patch('extraction.OpenAI', side_effect=self.new_client)
        self.factory.start()
        self.addCleanup(self.factory.stop)

    def new_client(self, **kwargs):
        return OpenAI(**kwargs, http_client=httpx.Client(transport=httpx.MockTransport(self.handle)))

    def handle(self, request):
        self.calls.append(json.loads(request.content))
        self.assertEqual(str(request.url), 'https://api.openai.com/v1/responses')
        if self.status_code != 200:
            return httpx.Response(self.status_code, json={'error': {'message': 'sensitive upstream error', 'type': 'api_error', 'code': None}})
        content = {'type': 'refusal', 'refusal': 'Cannot comply'} if self.refusal else {
            'type': 'output_text', 'annotations': [],
            'text': json.dumps({'fields': self.fields, 'notes': ['Verify unclear entries.']})}
        return httpx.Response(200, json={
            'id': 'resp_test', 'object': 'response', 'created_at': 1,
            'model': 'gpt-4.1-mini', 'status': self.response_status,
            'output': [{'type': 'message', 'id': 'msg_test', 'role': 'assistant', 'status': 'completed', 'content': [content]}],
        })

    def upload(self, data=None, filename='form.png'):
        return self.client.post('/api/extract', data={'files': (io.BytesIO(data if data is not None else image_bytes()), filename)})

    def test_image_uses_real_sdk_structured_response(self):
        response = self.upload(filename='original form.png')
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json['errors'], [])
        row = response.json['results'][0]
        self.assertEqual(row['filename'], 'original form.png')
        self.assertEqual(row['fields'], {'Name': 'Alice Smith', 'Phone': ''})
        call = self.calls[0]
        self.assertFalse(call['store'])
        self.assertEqual(call['text']['format']['type'], 'json_schema')
        self.assertTrue(call['text']['format']['strict'])
        content = call['input'][1]['content']
        image = next(p for p in content if p['type'] == 'input_image')
        self.assertEqual(image['detail'], 'high')
        self.assertTrue(base64.b64decode(image['image_url'].split(',')[1]).startswith(b'\x89PNG'))

    def test_all_pdf_pages_sent_as_images(self):
        with pymupdf.open() as document:
            document.new_page().insert_text((72, 72), 'Name: Alice')
            document.new_page().insert_image(pymupdf.Rect(0, 0, 400, 200), stream=image_bytes())
            response = self.upload(document.tobytes(), 'digital-and-scanned.pdf')
        self.assertEqual(response.json['errors'], [])
        self.assertEqual(len(self.calls), 1)
        content = self.calls[0]['input'][1]['content']
        overviews = [p['text'] for p in content if p['type'] == 'input_text' and 'overview' in p['text']]
        self.assertEqual(overviews, ['Page 1: full page overview', 'Page 2: full page overview'])
        images = [p for p in content if p['type'] == 'input_image']
        self.assertEqual(len(images), 10)  # Each rendered page has an overview and four close-ups.
        with Image.open(io.BytesIO(base64.b64decode(images[1]['image_url'].split(',')[1]))) as crop:
            # Detail views keep a higher pixel density than the old 2.5x PDF render.
            self.assertGreater(crop.height, 1800)

    def test_form_column_order_survives_json_and_exports(self):
        self.fields = [{'label': label, 'value': value} for label, value in
                       [('Z first field', '001'), ('A second field', ''), ('M third field', 'last')]]
        fields = self.upload().json['results'][0]['fields']
        labels = ['Z first field', 'A second field', 'M third field']
        self.assertEqual(list(fields), labels)
        body = {'columns': ['File name', *labels],
                'rows': [['form.png', *fields.values()]]}
        csv_result = self.client.post('/api/export/csv', json=body)
        self.assertEqual(next(csv.reader(io.StringIO(csv_result.data.decode('utf-8-sig')))), body['columns'])
        xlsx_result = self.client.post('/api/export/xlsx', json=body)
        workbook = load_workbook(io.BytesIO(xlsx_result.data))
        self.assertEqual([cell.value for cell in workbook.active[1]], body['columns'])

    def test_luna_default_is_sent_to_openai(self):
        with patch.dict(os.environ, {'OPENAI_MODEL': ''}):
            self.assertEqual(self.upload().json['errors'], [])
        self.assertEqual(self.calls[0]['model'], 'gpt-5.6-luna')

    def test_general_forms_keep_only_their_own_fields(self):
        layouts = [
            [('Student name', 'Asha'), ('Roll number', '0007'), ('Guardian name', '')],
            [('Teacher name', 'Ravi'), ('Subject', 'Math'), ('Employee ID', '0012')],
            [('Employee name', 'Meera'), ('Department', 'Accounts'), ('Manager name', '')],
            [('Applicant name', 'Asha'), ('Name', 'Self'), ('Address', '')],
        ]
        for layout in layouts:
            with self.subTest(layout=layout):
                self.fields = [{'label': label, 'value': value} for label, value in layout]
                fields = self.upload().json['results'][0]['fields']
                self.assertEqual(list(fields.items()), layout)
        prompt = self.calls[-1]['input'][0]['content']
        self.assertIn('Do not impose a predefined template', prompt)
        self.assertNotIn('Purnea', prompt)
        properties = self.calls[-1]['text']['format']['schema']['properties']
        self.assertEqual(set(properties), {'fields', 'notes'})

    def test_detail_views_keep_edges_and_overlap_without_lossy_encoding(self):
        source = Image.new('RGB', (2000, 2000), 'white')
        for x, y, color in [(0, 0, 'red'), (1000, 0, 'green'), (0, 1000, 'blue'), (1000, 1000, 'yellow')]:
            source.paste(color, (x, y, x + 1000, y + 1000))
        # Thin strokes at the center must survive in all overlapping close-ups.
        source.paste('black', (998, 998, 1002, 1002))
        output = io.BytesIO()
        source.save(output, 'PNG')
        content = document_content(output.getvalue(), 'small-text.png')
        images = [Image.open(io.BytesIO(base64.b64decode(p['image_url'].split(',')[1])))
                  for p in content if p['type'] == 'input_image']
        self.assertEqual(len(images), 5)
        self.assertEqual(images[0].size, source.size)
        corners = [(0, 0), (1099, 0), (0, 1099), (1099, 1099)]
        colors = [(255, 0, 0), (0, 128, 0), (0, 0, 255), (255, 255, 0)]
        centers = [(1000, 1000), (100, 1000), (1000, 100), (100, 100)]
        for crop, corner, color, center in zip(images[1:], corners, colors, centers):
            self.assertEqual(crop.getpixel(corner), color)
            self.assertEqual(crop.getpixel(center), (0, 0, 0))
        for image in images:
            image.close()

    def test_image_orientation_and_transparency(self):
        source = Image.new('RGBA', (100, 200), (0, 0, 0, 0))
        source.putpixel((0, 0), (255, 0, 0, 255))
        exif = Image.Exif()
        exif[274] = 6  # Rotate clockwise into the intended reading orientation.
        output = io.BytesIO()
        source.save(output, 'PNG', exif=exif)
        content = document_content(output.getvalue(), 'rotated.png')
        with Image.open(io.BytesIO(base64.b64decode(content[1]['image_url'].split(',')[1]))) as image:
            self.assertEqual(image.size, (200, 100))
            self.assertEqual(image.getpixel((199, 0)), (255, 0, 0))
            self.assertEqual(image.getpixel((0, 0)), (255, 255, 255))

    def test_rendered_size_limit_stops_before_api_call(self):
        with patch('extraction.MAX_ENCODED_BYTES', 10):
            self.assertIn('Split', self.upload().json['errors'][0]['error'])
        self.assertEqual(self.calls, [])

    def test_multiframe_tiff_preserves_page_order(self):
        output = io.BytesIO()
        Image.new('RGB', (100, 200), 'red').save(
            output, 'TIFF', save_all=True, append_images=[Image.new('RGB', (100, 200), 'blue')])
        content = document_content(output.getvalue(), 'two-pages.tiff')
        self.assertEqual([p['text'] for p in content if p['type'] == 'input_text'],
                         ['Page 1: full page overview', 'Page 2: full page overview'])
        images = [p for p in content if p['type'] == 'input_image']
        for part, color in zip(images, [(255, 0, 0), (0, 0, 255)]):
            with Image.open(io.BytesIO(base64.b64decode(part['image_url'].split(',')[1]))) as image:
                self.assertEqual(image.getpixel((0, 0)), color)

    def test_image_formats(self):
        for format, extension in [('BMP', 'bmp'), ('TIFF', 'tiff'), ('WEBP', 'webp'), ('JPEG', 'jpg')]:
            with self.subTest(format=format):
                self.assertEqual(self.upload(image_bytes(format), f'form.{extension}').json['errors'], [])

    def test_missing_key_and_key_not_in_html(self):
        page = self.client.get('/').data.decode()
        self.assertIn('OpenAI key configured', page)
        self.assertNotIn('test-secret-not-a-real-key', page)
        with patch.dict(os.environ, {'OPENAI_API_KEY': ''}):
            self.assertEqual(self.upload().status_code, 503)
            self.assertIn('Set up your OpenAI API key', self.client.get('/').data.decode())
        self.assertEqual(self.calls, [])

    def test_api_errors_are_actionable_and_redacted(self):
        for code, expected in [(401, 'API key'), (403, 'permission'), (404, 'model'), (429, 'quota'), (500, 'could not process')]:
            with self.subTest(code=code):
                self.status_code = code
                response = self.upload().json
                self.assertEqual(response['results'], [])
                self.assertIn(expected, response['errors'][0]['error'])
                self.assertNotIn('sensitive upstream', str(response))

    def test_refusal_and_incomplete_are_not_successes(self):
        self.refusal = True
        self.assertIn('declined', self.upload().json['errors'][0]['error'])
        self.refusal = False
        self.response_status = 'incomplete'
        self.assertIn('did not finish', self.upload().json['errors'][0]['error'])

    def test_duplicate_and_reserved_labels(self):
        self.fields = [{'label': 'File name', 'value': 'printed-value'}, {'label': 'Name', 'value': 'Alice'}, {'label': 'Name', 'value': 'Bob'}]
        fields = self.upload().json['results'][0]['fields']
        self.assertEqual(fields, {'File name (form field)': 'printed-value', 'Name': 'Alice', 'Name (2)': 'Bob'})

    def test_addresses_are_single_line_and_identifiers_stay_exact(self):
        self.fields = [
            {'label': ' Address :', 'value': '  Flat 04,\r\n12 Main Road\n\tPune  001234  '},
            {'label': 'Permanent ADDRESS', 'value': 'PO Box 007\u2028Delhi\u00a0110001'},
            {'label': 'Address', 'value': 'House 2\rMumbai'},
            {'label': 'Addr.', 'value': 'Line A\nLine B'},
            {'label': 'Address1', 'value': 'Road 1\nBlock 2'},
            {'label': 'Phone', 'value': '+91 00123-45678'},
            {'label': 'Account number', 'value': '00000120'},
            {'label': 'Remarks', 'value': 'First line\nSecond line'},
        ]
        fields = self.upload().json['results'][0]['fields']
        self.assertEqual(fields['Address'], 'Flat 04, 12 Main Road Pune 001234')
        self.assertEqual(fields['Permanent ADDRESS'], 'PO Box 007 Delhi 110001')
        self.assertEqual(fields['Address (2)'], 'House 2 Mumbai')
        self.assertEqual(fields['Addr.'], 'Line A Line B')
        self.assertEqual(fields['Address1'], 'Road 1 Block 2')
        self.assertEqual(fields['Phone'], '+91 00123-45678')
        self.assertEqual(fields['Account number'], '00000120')
        self.assertEqual(fields['Remarks'], 'First line\nSecond line')

    def test_invalid_files_and_page_limit_never_call_api(self):
        for data, filename in [(b'bad', 'bad.pdf'), (b'bad', 'bad.png'), (b'bad', 'bad.exe'), (b'', 'empty.pdf')]:
            self.assertTrue(self.upload(data, filename).json['errors'])
        with pymupdf.open() as document:
            for _ in range(21):
                document.new_page()
            self.assertTrue(self.upload(document.tobytes(), 'long.pdf').json['errors'])
        self.assertEqual(self.calls, [])

    def test_batch_errors_preserve_successful_files(self):
        response = self.client.post('/api/extract', data={'files': [(io.BytesIO(b'bad'), 'bad.pdf'), (io.BytesIO(image_bytes()), 'good.png')]})
        self.assertEqual(len(response.json['results']), 1)
        self.assertEqual(len(response.json['errors']), 1)

    def test_thousand_file_upload_and_export_boundary(self):
        # Exercise multipart parsing and output at the limit without 1,000 API calls.
        def fake_extract(data, filename):
            return {'filename': filename, 'fields': {'Name': 'Alice'}, 'text': '', 'method': 'test', 'warning': ''}
        with patch('app.extract', side_effect=fake_extract) as mocked:
            response = self.client.post('/api/extract', data={'files': [
                (io.BytesIO(b'form'), f'form-{i}.pdf') for i in range(1000)]})
            self.assertEqual(response.status_code, 200)
            self.assertEqual(len(response.json['results']), 1000)
            self.assertEqual(mocked.call_count, 1000)
            body = {'columns': ['File name', 'Name'], 'rows': [
                [row['filename'], row['fields']['Name']] for row in response.json['results']]}
            export = self.client.post('/api/export/csv', json=body)
            rows = list(csv.reader(io.StringIO(export.data.decode('utf-8-sig'))))
            self.assertEqual(len(rows), 1001)
            self.assertEqual(rows[-1], ['form-999.pdf', 'Alice'])
            mocked.reset_mock()
            response = self.client.post('/api/extract', data={'files': [
                (io.BytesIO(b'form'), f'form-{i}.pdf') for i in range(1001)]})
            self.assertEqual(response.status_code, 400)
            self.assertIn('1,000', response.json['error'])
            mocked.assert_not_called()

    def test_exports_preserve_edits_and_prevent_formulas(self):
        body = {'columns': ['File name', 'Name', 'Phone'], 'rows': [['form.pdf', 'Alice, Smith', '+12345'], ['second.png', '=1+1', '']]}
        response = self.client.post('/api/export/csv', json=body)
        rows = list(csv.reader(io.StringIO(response.data.decode('utf-8-sig'))))
        self.assertEqual(rows[1], ['form.pdf', 'Alice, Smith', "'+12345"])
        self.assertEqual(rows[2][1], "'=1+1")
        response = self.client.post('/api/export/xlsx', json=body)
        workbook = load_workbook(io.BytesIO(response.data))
        self.assertEqual(workbook.active['A2'].value, 'form.pdf')
        self.assertEqual(workbook.active['B3'].data_type, 's')

    def test_exports_flatten_edited_addresses_and_keep_formula_protection(self):
        body = {'columns': ['File name', 'Postal Address', 'Address (2)', 'Account number', 'Remarks'],
                'rows': [['form.pdf', '  001 Main St\r\nFlat\t04\u2028Pune 411001  ',
                          '\n =1+1', '00000120', 'First line\nSecond line']]}
        expected = ['form.pdf', '001 Main St Flat 04 Pune 411001', "'=1+1", '00000120', 'First line\nSecond line']
        response = self.client.post('/api/export/csv', json=body)
        rows = list(csv.reader(io.StringIO(response.data.decode('utf-8-sig'))))
        self.assertEqual(rows[1], expected)
        response = self.client.post('/api/export/xlsx', json=body)
        with io.BytesIO(response.data) as output:
            workbook = load_workbook(output)
            self.assertEqual([cell.value for cell in workbook.active[2]], expected)
            self.assertEqual(workbook.active['C2'].data_type, 's')
            self.assertFalse(workbook.active['B2'].alignment.wrap_text)

    def test_health_and_validation(self):
        self.assertEqual(self.client.get('/health').json, {'status': 'ok'})
        self.assertEqual(self.client.post('/api/extract').status_code, 400)
        self.assertEqual(self.client.post('/api/export/xlsx', json={'columns': ['Name'], 'rows': [[]]}).status_code, 400)


if __name__ == '__main__':
    unittest.main(verbosity=2)
