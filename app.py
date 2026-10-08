import csv
import io

from flask import Flask, jsonify, render_template, request, send_file
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from werkzeug.exceptions import RequestEntityTooLarge

from extraction import ExtractionError, configuration, extract, normalize_field_value

app = Flask(__name__)
# Preserve the model's form order when serializing fields for the browser.
app.json.sort_keys = False
app.config['MAX_CONTENT_LENGTH'] = 50 * 1024 * 1024
app.config['MAX_FORM_PARTS'] = 1001


@app.get('/')
def index():
    return render_template('index.html', openai=configuration())


@app.get('/health')
def health():
    return jsonify(status='ok')


@app.errorhandler(RequestEntityTooLarge)
def too_large(error):
    return jsonify(error='Upload exceeds the 50 MB request limit or the allowed number of form parts.'), 413


@app.post('/api/extract')
def extract_files():
    files = request.files.getlist('files')
    if not files or len(files) > 1000:
        return jsonify(error='Select between 1 and 1,000 files.'), 400
    if not configuration()['configured']:
        return jsonify(error='Set OPENAI_API_KEY in .env and run docker compose up -d to enable extraction.'), 503
    results, errors = [], []
    for upload in files:
        filename = (upload.filename or 'untitled').replace('\\', '/').rsplit('/', 1)[-1]
        try:
            data = upload.read()
            if not data:
                raise ValueError('The file is empty.')
            results.append(extract(data, filename))
        except (ValueError, ExtractionError) as error:
            errors.append({'filename': filename, 'error': str(error)})
        except Exception:
            app.logger.error('Document extraction failed unexpectedly')
            errors.append({'filename': filename, 'error': 'Unable to extract this document. Try a clearer or smaller file.'})
    return jsonify(results=results, errors=errors)


def safe_cell(value):
    value = ILLEGAL_CHARACTERS_RE.sub('', str(value if value is not None else ''))
    # Prevent spreadsheet formulas supplied by documents or edited cells.
    if value.lstrip().startswith(('=', '+', '-', '@')) or value.startswith(('\t', '\r', '\n')):
        value = "'" + value
    return value


@app.post('/api/export/<kind>')
def export(kind):
    if kind not in {'csv', 'xlsx'}:
        return jsonify(error='Choose CSV or XLSX.'), 400
    body = request.get_json(silent=True) or {}
    columns, rows = body.get('columns'), body.get('rows')
    if not isinstance(columns, list) or not columns or len(columns) > 500 or not all(isinstance(c, str) for c in columns):
        return jsonify(error='Invalid columns (maximum 500).'), 400
    if not isinstance(rows, list) or len(rows) > 1000 or not all(isinstance(r, list) and len(r) == len(columns) for r in rows):
        return jsonify(error='Invalid rows (maximum 1,000).'), 400
    values = [[safe_cell(c) for c in columns]] + [
        [safe_cell(normalize_field_value(column, value)) for column, value in zip(columns, row)]
        for row in rows]
    output = io.BytesIO()
    if kind == 'csv':
        stream = io.StringIO(newline='')
        csv.writer(stream).writerows(values)
        output.write(stream.getvalue().encode('utf-8-sig'))
        mimetype = 'text/csv'
    else:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = 'Form data'
        for row in values:
            sheet.append(row)
        sheet.freeze_panes = 'B2'
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            from openpyxl.styles import Font, PatternFill
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='146C60')
            sheet.column_dimensions[cell.column_letter].width = 26
        workbook.save(output)
        mimetype = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    output.seek(0)
    return send_file(output, mimetype=mimetype, as_attachment=True, download_name=f'form-data.{kind}')


if __name__ == '__main__':
    app.run(host='0.0.0.0', port=8086)
